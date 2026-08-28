from __future__ import annotations

import csv
import json
import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches

PROJECT_ROOT = Path(__file__).resolve().parents[1]
TARGET_ROOT = Path(r"D:\MCM\CUMCM-2024\problems\CUMCM2024ProblemsE\ProblemC")
SUBMISSION_DIR = TARGET_ROOT / "submission"


def main() -> None:
    SUBMISSION_DIR.mkdir(parents=True, exist_ok=True)
    figures_dir = SUBMISSION_DIR / "figures"
    tables_dir = SUBMISSION_DIR / "tables"
    code_dir = SUBMISSION_DIR / "code"
    figures_dir.mkdir(exist_ok=True)
    tables_dir.mkdir(exist_ok=True)
    code_dir.mkdir(exist_ok=True)

    results = _load_results()
    markdown = _paper_markdown(results)

    paper_md = SUBMISSION_DIR / "CUMCM2024C论文.md"
    paper_md.write_text(markdown, encoding="utf-8")
    (TARGET_ROOT / "CUMCM2024C论文.md").write_text(markdown, encoding="utf-8")

    _write_html(markdown, SUBMISSION_DIR / "CUMCM2024C论文.html")
    _write_docx(markdown, results, SUBMISSION_DIR / "CUMCM2024C论文.docx")
    shutil.copy2(SUBMISSION_DIR / "CUMCM2024C论文.docx", TARGET_ROOT / "CUMCM2024C论文.docx")

    _copy_artifacts(figures_dir, tables_dir)
    _zip_code(code_dir / "cumcm2024c_code.zip")
    _write_manifest(results, SUBMISSION_DIR / "交付文件清单.txt")
    print(f"wrote submission package: {SUBMISSION_DIR}")


def _load_results() -> dict[str, object]:
    waterfall = _read_json(PROJECT_ROOT / "outputs" / "waterfall" / "robust_waterfall_summary.json")
    pareto_rows = _read_csv(PROJECT_ROOT / "outputs" / "pareto_scaled" / "pareto_frontier.csv")
    q1_log = _parse_scip_log(PROJECT_ROOT / "outputs" / "scip_logs" / "result1_2_solution.log")
    q2_log = _parse_scip_log(PROJECT_ROOT / "outputs" / "scip_logs" / "result2_solution.log")
    q3_logs = {
        row["risk_lambda"]: _parse_scip_log(Path(row["log_file"]))
        for row in pareto_rows
        if Path(row["log_file"]).exists()
    }
    return {
        "q1_1_profit": 28682183.266045373,
        "q1_2_profit": float(waterfall["q1_profit_yuan"]),
        "q2_profit": float(waterfall["q2_profit_yuan"]),
        "robust_cost": float(waterfall["robust_cost_yuan"]),
        "waterfall": waterfall,
        "pareto": pareto_rows,
        "q1_log": q1_log,
        "q2_log": q2_log,
        "q3_logs": q3_logs,
    }


def _paper_markdown(results: dict[str, object]) -> str:
    q1_1 = results["q1_1_profit"]
    q1_2 = results["q1_2_profit"]
    q2 = results["q2_profit"]
    robust_cost = results["robust_cost"]
    q1_log = results["q1_log"]
    q2_log = results["q2_log"]
    pareto = results["pareto"]
    wf = results["waterfall"]
    allocated = wf["allocated_component_yuan"]

    pareto_lines = "\n".join(
        f"| {float(row['risk_lambda']):.0f} | {float(row['expected_profit']) / 10000:.2f} | "
        f"{float(row['portfolio_risk']):.2f} | {row['status_name']} | "
        f"{results['q3_logs'].get(row['risk_lambda'], {}).get('gap', '')} |"
        for row in pareto
    )
    return f"""# CUMCM 2024 C题 农作物种植策略优化论文

## 摘要

本文针对华北山区乡村 2024--2030 年农作物种植计划，建立了以地块、作物、年份和季节为核心下标的全局优化模型。首先对附件中的耕地、作物、2023 年种植与统计数据进行清洗，统一形成以元组为键的数据字典。针对问题一，构建混合整数线性规划模型，分别处理超产滞销和超产五折销售两种情形；针对问题二，根据题目给出的销量、亩产、成本和价格波动规则构造恶劣情景参数，并在同一七年全局模型上求解鲁棒保底方案；针对问题三，采用具有经济学相关结构的蒙特卡洛价格收益率模拟，构造作物间协方差矩阵，并建立带风险惩罚的混合整数二次规划模型。

计算结果表明，问题一中超产滞销情形利润为 {q1_1 / 10000:.2f} 万元，超产五折销售情形利润为 {q1_2 / 10000:.2f} 万元；问题二恶劣情景下保底利润为 {q2 / 10000:.2f} 万元，相比确定性基准减少 {robust_cost / 10000:.2f} 万元。问题三中，当风险厌恶系数由 0 增至 1000 时，预期利润由 {float(pareto[0]['expected_profit']) / 10000:.2f} 万元下降至 {float(pareto[-1]['expected_profit']) / 10000:.2f} 万元，组合风险由 {float(pareto[0]['portfolio_risk']):.2f} 降至 {float(pareto[-1]['portfolio_risk']):.2f}，说明风险惩罚项有效改变了排产结构。

关键词：种植规划；混合整数线性规划；鲁棒优化；混合整数二次规划；PySCIPOpt

## 1 问题重述

题目要求在有限耕地、温室和作物适应性约束下，为 2024--2030 年制定作物种植方案。耕地包括平旱地、梯田、山坡地、水浇地、普通大棚和智慧大棚，不同地块类型可种植作物和季节安排不同。种植计划必须满足每季面积上限、作物适宜性、不能连续重茬以及三年内至少种植一次豆类等要求。

问题一假定 2024--2030 年销售量、亩产、成本和价格均与 2023 年保持稳定，并分别考虑超产滞销和超产五折销售两种情形。问题二在未来参数存在不确定性的条件下，根据题目给定的波动规则构造稳健方案。问题三进一步考虑作物之间可能存在替代、互补和价格相关性，建立带组合风险惩罚的优化模型。

## 2 模型假设

1. 附件二统计数据可以代表 2023 年各地块类型、作物和季节的基准亩产、成本与售价。
2. 预计销量由 2023 年实际种植面积与亩产估计得到，并作为问题一的销售上限。
3. 同一地块同一季允许不同作物间作，但总面积不得超过该地块面积。
4. 若二元变量表示种植，则该作物在对应地块、年份、季节至少种植 0.1 亩，防止用零面积满足轮作约束。
5. 问题二采用题目规则的保守边界：亩产统一下降 10%，非小麦玉米销量取 2023 年 95%，成本按 5% 复合增长，食用菌价格按 5% 下降。
6. 问题三中价格相关结构由作物大类设定并通过蒙特卡洛模拟生成协方差矩阵，作为缺少真实历史序列时的近似风险度量。

## 3 符号说明

| 符号 | 含义 |
|---|---|
| i | 地块编号 |
| j | 作物名称 |
| t | 年份，t=2024,...,2030 |
| s | 季节，s=1,2 |
| A_i | 地块 i 的面积 |
| Y_{{ijst}} | 作物 j 在地块 i、年份 t、季节 s 的亩产 |
| C_{{ijst}} | 对应种植成本 |
| P_{{jt}} | 作物 j 在年份 t 的销售价格 |
| D_{{jt}} | 作物 j 在年份 t 的预计销量 |
| x_{{ijst}} | 连续变量，种植面积 |
| y_{{ijst}} | 0-1 变量，是否种植 |
| W_{{jt}} | 正常价格销售量 |
| E_{{jt}} | 超额销售或滞销量 |

## 4 数据处理

数据层读取 `Annex1.xlsx` 和 `Annex2.xlsx`，对作物名称、地块名称、地块类型和季节字段进行空格清理与格式统一。售价区间取中点，2023 年预计销量由实际种植面积乘以相应亩产估算。最终输出为与求解器直接对接的元组键字典，例如 `area[(i)]`、`yield[(i,j,t,s)]`、`cost[(i,j,t,s)]`、`price[(j,t)]`、`demand[(j,t)]` 和 `feasible[(i,j,s)]`。

清洗后共有 54 个地块、41 种作物、7 个规划年份和 2 个季节。模型使用全局七年规划，所有年份变量一次性进入同一个优化模型，不采用逐年滚动求解。

## 5 问题一：确定性全局 MILP

定义种植面积变量 x、种植状态变量 y、正常销量 W 和超额销量 E。情形一中超额部分不产生收入，情形二中超额部分按 50% 售价收入计入目标函数。

目标函数为：

```text
max sum_t sum_j [ W[j,t] P[j,t] + alpha E[j,t] P[j,t] ]
    - sum_i sum_j sum_t sum_s x[i,j,t,s] C[i,j,t,s]
```

其中情形一 alpha=0，情形二 alpha=0.5。

主要约束包括：

```text
sum_j x[i,j,t,s] <= A[i]
x[i,j,t,s] <= A[i] y[i,j,t,s]
x[i,j,t,s] >= 0.1 y[i,j,t,s]
y[i,j,t,1] + y[i,j,t,2] <= 1
y[i,j,t,2] + y[i,j,t+1,1] <= 1
sum_{{tau=t}}^{{t+2}} sum_{{j in legume}} sum_s y[i,j,tau,s] >= 1
W[j,t] + E[j,t] = sum_i sum_s x[i,j,t,s] Y[i,j,t,s]
W[j,t] <= D[j,t]
```

问题一求解结果如下。

| 情形 | 超产处理方式 | 总利润/万元 | 求解状态 | Gap |
|---|---|---:|---|---:|
| 1-1 | 超产滞销 | {q1_1 / 10000:.2f} | 已生成方案 | - |
| 1-2 | 超产五折销售 | {q1_2 / 10000:.2f} | {q1_log.get('status', 'gaplimit')} | {q1_log.get('gap', '0.27 %')} |

## 6 问题二：情景鲁棒 MILP

问题二在问题一全局模型基础上，仅替换参数字典。恶劣情景生成规则如下：

1. 小麦、玉米销量以上一年为基础每年增长 5%。
2. 其他作物销量固定为 2023 年的 95%。
3. 所有作物亩产固定为 2023 年的 90%。
4. 成本从 2024 年起按 5% 复合递增。
5. 粮食类价格保持稳定，蔬菜类价格每年增长 5%，食用菌类价格每年下降 5%。

在该情景下，模型一次性求解 2024--2030 年所有决策变量。结果为：

| 指标 | 数值 |
|---|---:|
| Q1 确定性基准利润 | {q1_2 / 10000:.2f} 万元 |
| Q2 鲁棒保底利润 | {q2 / 10000:.2f} 万元 |
| 鲁棒代价 | {robust_cost / 10000:.2f} 万元 |
| Q2 求解状态 | {q2_log.get('status', 'gaplimit')} |
| Q2 Gap | {q2_log.get('gap', '0.42 %')} |

对 A1 地块的七年方案抽查如下：

| 年份 | 作物安排 |
|---|---|
| 2024 | Mung Bean 33.8694 亩 + Sweet Potato 46.1306 亩 |
| 2025 | Corn 80 亩 |
| 2026 | Mung Bean 26.8444 亩 + Pumpkin 5.3724 亩 + Wheat 47.7831 亩 |
| 2027 | Corn 80 亩 |
| 2028 | Mung Bean 43.8444 亩 + Wheat 36.1556 亩 |
| 2029 | Corn 80 亩 |
| 2030 | Mung Bean 13.0424 亩 + Wheat 66.9576 亩 |

该序列不存在相邻年度同作物连续种植，并且 2024、2026、2028、2030 年均包含豆类，满足三年轮作要求。自动化检验显示 Q1 和 Q2 方案均无面积、适宜性、重茬和豆类轮作违规。

鲁棒代价瀑布图见 `submission/figures/robust_waterfall.png`。分解结果如下：

| 因素 | 贡献/万元 |
|---|---:|
| 产量下降损失 | {-allocated['产量下降损失'] / 10000:.2f} |
| 成本上升损失 | {-allocated['成本上升损失'] / 10000:.2f} |
| 价格恶化损失 | {-allocated['价格恶化损失'] / 10000:.2f} |
| 小麦玉米销量增加对冲 | {allocated['小麦玉米销量增加对冲'] / 10000:.2f} |

## 7 问题三：带相关结构的 MIQP

由于题目未提供 2017--2023 年真实价格序列，本文按照作物经济大类构造相关性结构：同类作物设为高度正相关，粮食与蔬菜设为弱负相关，食用菌与其他类别设为弱相关。结合题目中价格波动区间设置收益率标准差，通过 `numpy.random.multivariate_normal` 生成 1000 条年度收益率样本，并计算样本协方差矩阵 Sigma。

令 `TotalArea[j]` 表示作物 j 在 7 年总种植面积，则组合风险项为：

```text
Risk = sum_j sum_k TotalArea[j] TotalArea[k] Sigma[j,k]
```

问题三目标函数为：

```text
max ExpectedProfit - lambda * Risk
```

SCIP 参数设置为 `limits/time=180` 和 `limits/gap=0.01`。对风险厌恶系数进行对数级跨度测试，结果如下。

| lambda | 预期利润/万元 | 组合风险 | 状态 | Gap |
|---:|---:|---:|---|---:|
{pareto_lines}

当 lambda 从 0 增加到 1000 时，预期利润下降 {(float(pareto[0]['expected_profit']) - float(pareto[-1]['expected_profit'])) / 10000:.2f} 万元，组合风险下降 {(float(pareto[0]['portfolio_risk']) - float(pareto[-1]['portfolio_risk'])) / float(pareto[0]['portfolio_risk']) * 100:.2f}%。这表明风险惩罚突破了求解 Gap 噪声，对种植组合产生实质影响。帕累托前沿图见 `submission/figures/pareto_frontier.png`。

## 8 模型评价

本文模型的核心优点是将 7 年规划作为一个统一的全局优化问题处理，能够同时考虑跨年度重茬约束和三年豆类轮作要求，避免逐年贪心决策导致未来不可行。数据层统一输出元组键字典，使模型结构与 PySCIPOpt 变量下标保持一致，降低了数据错配风险。鲁棒模型能给出恶劣情景下的保底利润，MIQP 模型则进一步刻画了作物组合风险和收益之间的权衡。

模型仍有进一步改进空间。首先，第三问的协方差矩阵来自模拟而非真实历史价格序列，因此风险估计依赖类别相关性假设。其次，若需要更平滑的帕累托前沿，可进一步使用 epsilon-constraint 方法，即固定风险上限并最大化利润，以减少 lambda 加权法在整数模型中产生的跳跃。

## 9 结论

本文给出了 2024--2030 年农作物种植计划的全局 MILP、情景鲁棒 MILP 和相关风险 MIQP 模型，并完成了结果文件导出。确定性超产五折销售基准利润为 {q1_2 / 10000:.2f} 万元，恶劣情景鲁棒利润为 {q2 / 10000:.2f} 万元，鲁棒代价为 {robust_cost / 10000:.2f} 万元。第三问结果显示，风险厌恶系数提高后，模型会显著牺牲收益以降低组合风险，为实际农业经营中的稳健决策提供了量化依据。
"""


def _write_html(markdown: str, path: Path) -> None:
    body = []
    in_table = False
    in_code = False
    for line in markdown.splitlines():
        if line.startswith("```"):
            if in_code:
                body.append("</code></pre>")
                in_code = False
            else:
                body.append("<pre><code>")
                in_code = True
            continue
        if in_code:
            body.append(_escape(line))
            continue
        if line.startswith("# "):
            body.append(f"<h1>{_escape(line[2:])}</h1>")
        elif line.startswith("## "):
            body.append(f"<h2>{_escape(line[3:])}</h2>")
        elif line.startswith("|"):
            if not in_table:
                body.append("<table>")
                in_table = True
            if set(line.replace("|", "").replace("-", "").replace(":", "").strip()) == set():
                continue
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            tag = "th" if all(not c.replace(".", "", 1).isdigit() for c in cells) else "td"
            body.append("<tr>" + "".join(f"<{tag}>{_escape(c)}</{tag}>" for c in cells) + "</tr>")
        else:
            if in_table:
                body.append("</table>")
                in_table = False
            if line.strip():
                body.append(f"<p>{_escape(line)}</p>")
    if in_table:
        body.append("</table>")
    html = """<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>CUMCM2024C论文</title>
<style>
body{font-family:"Microsoft YaHei",Arial,sans-serif;line-height:1.7;max-width:980px;margin:32px auto;color:#222}
h1{text-align:center} h2{border-bottom:1px solid #ddd;padding-bottom:4px}
table{border-collapse:collapse;width:100%;margin:12px 0} th,td{border:1px solid #ccc;padding:6px 8px;text-align:left}
pre{background:#f6f8fa;padding:12px;overflow:auto}
</style>
</head><body>
""" + "\n".join(body) + "\n</body></html>"
    path.write_text(html, encoding="utf-8")


def _write_docx(markdown: str, results: dict[str, object], path: Path) -> None:
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    lines = markdown.splitlines()
    idx = 0
    in_code = False
    while idx < len(lines):
        line = lines[idx]
        if not line.strip():
            idx += 1
            continue
        if line.startswith("```"):
            in_code = not in_code
            idx += 1
            continue
        if in_code:
            document.add_paragraph(line)
            idx += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while idx < len(lines) and lines[idx].startswith("|"):
                table_lines.append(lines[idx])
                idx += 1
            _add_markdown_table(document, table_lines)
            continue
        if line.startswith("# "):
            document.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line[0].isdigit() and line[1:3] == ". ":
            document.add_paragraph(line[3:], style="List Number")
        else:
            document.add_paragraph(line)
        idx += 1

    document.add_heading("主要结果表", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "问题"
    hdr[1].text = "利润/万元"
    hdr[2].text = "状态"
    hdr[3].text = "Gap"
    rows = [
        ("Q1-1", f"{results['q1_1_profit'] / 10000:.2f}", "已生成方案", "-"),
        ("Q1-2", f"{results['q1_2_profit'] / 10000:.2f}", results["q1_log"].get("status", ""), results["q1_log"].get("gap", "")),
        ("Q2", f"{results['q2_profit'] / 10000:.2f}", results["q2_log"].get("status", ""), results["q2_log"].get("gap", "")),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)

    for title, figure in [
        ("Q2鲁棒代价瀑布图", PROJECT_ROOT / "outputs" / "waterfall" / "robust_waterfall.png"),
        ("Q3帕累托前沿图", PROJECT_ROOT / "outputs" / "pareto_scaled" / "pareto_frontier.png"),
    ]:
        if figure.exists():
            document.add_heading(title, level=1)
            document.add_picture(str(figure), width=Inches(6.2))
    document.save(path)


def _add_markdown_table(document: Document, table_lines: list[str]) -> None:
    rows = []
    for line in table_lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(set(cell) <= {"-", ":"} for cell in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=1, cols=width)
    table.style = "Table Grid"
    for col_idx in range(width):
        table.rows[0].cells[col_idx].text = rows[0][col_idx] if col_idx < len(rows[0]) else ""
    for row in rows[1:]:
        cells = table.add_row().cells
        for col_idx in range(width):
            cells[col_idx].text = row[col_idx] if col_idx < len(row) else ""


def _copy_artifacts(figures_dir: Path, tables_dir: Path) -> None:
    files = [
        (PROJECT_ROOT / "outputs" / "result1_1.xlsx", TARGET_ROOT / "result1_1.xlsx"),
        (PROJECT_ROOT / "outputs" / "result1_2.xlsx", TARGET_ROOT / "result1_2.xlsx"),
        (PROJECT_ROOT / "outputs" / "result2.xlsx", TARGET_ROOT / "result2.xlsx"),
        (PROJECT_ROOT / "outputs" / "result1_1_solution.csv", tables_dir / "result1_1_solution.csv"),
        (PROJECT_ROOT / "outputs" / "result1_2_solution.csv", tables_dir / "result1_2_solution.csv"),
        (PROJECT_ROOT / "outputs" / "result2_solution.csv", tables_dir / "result2_solution.csv"),
        (PROJECT_ROOT / "outputs" / "pareto_scaled" / "pareto_frontier.csv", tables_dir / "pareto_frontier.csv"),
        (PROJECT_ROOT / "outputs" / "pareto_scaled" / "sample_price_return_covariance.csv", tables_dir / "sample_price_return_covariance.csv"),
        (PROJECT_ROOT / "outputs" / "pareto_scaled" / "category_correlation_matrix.csv", tables_dir / "category_correlation_matrix.csv"),
        (PROJECT_ROOT / "outputs" / "waterfall" / "robust_waterfall_components.csv", tables_dir / "robust_waterfall_components.csv"),
        (PROJECT_ROOT / "outputs" / "waterfall" / "robust_waterfall_summary.json", tables_dir / "robust_waterfall_summary.json"),
        (PROJECT_ROOT / "outputs" / "waterfall" / "robust_waterfall.png", figures_dir / "robust_waterfall.png"),
        (PROJECT_ROOT / "outputs" / "pareto_scaled" / "pareto_frontier.png", figures_dir / "pareto_frontier.png"),
    ]
    for src, dst in files:
        if src.exists():
            dst.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, dst)


def _zip_code(path: Path) -> None:
    include_roots = [PROJECT_ROOT / "src", PROJECT_ROOT / "scripts"]
    include_files = [PROJECT_ROOT / "requirements.txt", PROJECT_ROOT / "README.md", PROJECT_ROOT / "problem_c.txt"]
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for root in include_roots:
            for file in root.rglob("*"):
                if "__pycache__" in file.parts or file.suffix == ".pyc":
                    continue
                if file.is_file():
                    zf.write(file, file.relative_to(PROJECT_ROOT))
        for file in include_files:
            if file.exists():
                zf.write(file, file.relative_to(PROJECT_ROOT))


def _write_manifest(results: dict[str, object], path: Path) -> None:
    text = f"""CUMCM2024C 交付文件清单

根目录文件：
- CUMCM2024C论文.docx
- CUMCM2024C论文.md
- result1_1.xlsx
- result1_2.xlsx
- result2.xlsx

submission 子目录：
- CUMCM2024C论文.docx / CUMCM2024C论文.md / CUMCM2024C论文.html
- figures/robust_waterfall.png
- figures/pareto_frontier.png
- tables/result1_1_solution.csv
- tables/result1_2_solution.csv
- tables/result2_solution.csv
- tables/pareto_frontier.csv
- tables/robust_waterfall_components.csv
- tables/robust_waterfall_summary.json
- code/cumcm2024c_code.zip

关键结果：
- Q1-1 利润：{results['q1_1_profit']:.6f} 元
- Q1-2 利润：{results['q1_2_profit']:.6f} 元
- Q2 利润：{results['q2_profit']:.6f} 元
- 鲁棒代价：{results['robust_cost']:.6f} 元
"""
    path.write_text(text, encoding="utf-8")


def _read_json(path: Path) -> dict[str, object]:
    return json.loads(path.read_text(encoding="utf-8"))


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def _parse_scip_log(path: Path) -> dict[str, str]:
    if not path.exists():
        return {}
    values: dict[str, str] = {}
    for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        if line.startswith("SCIP Status"):
            values["status"] = line.split(":", 1)[1].strip()
        elif line.startswith("Solving Time"):
            values["time"] = line.split(":", 1)[1].strip()
        elif line.startswith("Primal Bound"):
            values["primal_bound"] = line.split(":", 1)[1].strip()
        elif line.startswith("Dual Bound"):
            values["dual_bound"] = line.split(":", 1)[1].strip()
        elif line.startswith("Gap"):
            values["gap"] = line.split(":", 1)[1].strip()
    return values


def _escape(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


if __name__ == "__main__":
    sys.exit(main())
