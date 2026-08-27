from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .config import DEFAULT_DATA_DIR
from .data import clean_text
from .model import SolveResult


def write_solution_to_template(
    result: SolveResult,
    *,
    template_name: str,
    output_path: Path,
    data_dir: Path = DEFAULT_DATA_DIR,
) -> None:
    template_path = Path(data_dir) / "Annex3" / template_name
    if not template_path.exists():
        raise FileNotFoundError(f"Cannot find template: {template_path}")

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(template_path)
    by_year = _group_solution(result)
    for sheet_name in workbook.sheetnames:
        year = int(clean_text(sheet_name))
        worksheet = workbook[sheet_name]
        crop_col = _crop_columns(worksheet)
        season_plot_row = _season_plot_rows(worksheet)
        _clear_body(worksheet, crop_col, season_plot_row)
        for (plot, crop, season), area in by_year.get(year, {}).items():
            row = season_plot_row.get((season, plot))
            col = crop_col.get(crop)
            if row is not None and col is not None:
                worksheet.cell(row=row, column=col, value=round(area, 4))
    workbook.save(output_path)


def _group_solution(result: SolveResult) -> dict[int, dict[tuple[str, str, int], float]]:
    grouped: dict[int, dict[tuple[str, str, int], float]] = {}
    for (plot, crop, year, season), area in result.planted_area.items():
        grouped.setdefault(year, {})[(plot, crop, season)] = grouped.setdefault(year, {}).get((plot, crop, season), 0.0) + area
    return grouped


def _crop_columns(worksheet) -> dict[str, int]:
    columns = {}
    for cell in worksheet[1]:
        crop = clean_text(cell.value)
        if crop and crop != "Plot Name":
            columns[crop] = cell.column
    return columns


def _season_plot_rows(worksheet) -> dict[tuple[int, str], int]:
    rows = {}
    season = 1
    for row in range(2, worksheet.max_row + 1):
        marker = clean_text(worksheet.cell(row=row, column=1).value).lower()
        if "second" in marker:
            season = 2
        plot = clean_text(worksheet.cell(row=row, column=2).value)
        if plot:
            rows[(season, plot)] = row
    return rows


def _clear_body(worksheet, crop_col: dict[str, int], season_plot_row: dict[tuple[int, str], int]) -> None:
    for row in season_plot_row.values():
        for col in crop_col.values():
            worksheet.cell(row=row, column=col, value=None)
